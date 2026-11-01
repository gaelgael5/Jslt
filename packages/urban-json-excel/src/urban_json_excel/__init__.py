"""Plugin urban-json : lecture de classeurs Excel (``loadxls``, ``loadxlssheet``).

Port de ``ServicesXls`` + ``ExcelReader`` (.NET, ExcelDataReader). ``.xlsx`` via
openpyxl ; ``.xls`` (Excel 97-2003) via xlrd, avec l'extra ``urban-json-excel[xls]``.

* ``loadxls(fichier, fichier_de_configuration)`` — configuration lue dans un
  fichier JSON ;
* ``loadxlssheet(fichier, configuration)`` — configuration passée en objet.

Configuration (clés insensibles à la casse, comme la désérialisation .NET) ::

    { "FirstRowHasColumnName": true,  // 1re ligne = noms de colonnes
      "HeaderLine": 0,                // > 0 : ligne d'en-tête (données après elle)
      "Tablename": "Feuil1",          // feuille ; vide = l'unique feuille
      "AddRowNumber": true,           // ajoute "$row" (rang de la ligne de données)
      "ByPassEmptyObject": true,      // omet les lignes sans aucune valeur
      "Columns": [ { "Index": 0, "Name": "id", "Type": "integer",
                     "Trimmed": true, "BypassEmptyValue": true } ] }

Sans ``Columns`` : toutes les colonnes, noms assainis (accents retirés, espaces ->
``_``, doublons suffixés ``_renamed_<index>``), colonne numérique entièrement
entière typée ``integer``. Avec ``Columns`` : seules ces colonnes, typées.

Écarts volontaires vs .NET : feuille unique prise quand ``Tablename`` est vide (le
.NET échouait toujours, son test « table manquante » précédant ce cas) ; dates en
ISO 8601 ; erreurs explicites (``JsltError``). Chemins relatifs au répertoire
courant (pas de ``ResolveFile``).
"""

from __future__ import annotations

import json
import os
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from typing import Any

from urban_json import JsltError

__version__ = "0.1.1"

__all__ = ["execute_loadxls", "execute_loadxlssheet", "register"]


# --- configuration ------------------------------------------------------------------


def _get(config: dict, key: str, default: Any = None) -> Any:
    lowered = key.lower()
    for name, value in config.items():
        if str(name).lower() == lowered:
            return value
    return default


@dataclass
class _Column:
    index: int
    name: str
    type: str = "string"
    trimmed: bool = True
    bypass_empty: bool = True


_TYPE_ALIASES = {
    "string": "string", "system.string": "string", "text": "string",
    "integer": "integer", "int": "integer", "long": "integer", "int32": "integer", "int64": "integer",
    "system.int32": "integer", "system.int64": "integer",
    "decimal": "number", "double": "number", "float": "number", "number": "number",
    "system.double": "number", "system.decimal": "number", "system.single": "number",
    "boolean": "boolean", "bool": "boolean", "system.boolean": "boolean",
    "datetime": "datetime", "date": "datetime", "system.datetime": "datetime",
    "auto": "auto",
}


def _columns_from_config(config: dict) -> list[_Column] | None:
    raw = _get(config, "Columns")
    if not raw:
        return None
    columns = []
    for item in raw:
        if not isinstance(item, dict) or _get(item, "Index") is None or not _get(item, "Name"):
            raise JsltError(f"loadxls : colonne de configuration invalide {item!r} (Index et Name requis)")
        type_name = str(_get(item, "Type", "string")).strip().lower()
        if type_name not in _TYPE_ALIASES:
            raise JsltError(f"loadxls : type de colonne inconnu {type_name!r}")
        columns.append(_Column(
            index=int(_get(item, "Index")),
            name=str(_get(item, "Name")),
            type=_TYPE_ALIASES[type_name],
            trimmed=bool(_get(item, "Trimmed", True)),
            bypass_empty=bool(_get(item, "BypassEmptyValue", True)),
        ))
    return columns


# --- lecture du classeur ---------------------------------------------------------------


def _read_sheets(path: str) -> dict[str, list[tuple]]:
    if path.lower().endswith(".xls"):
        try:
            import xlrd
        except ImportError as exc:
            raise JsltError("loadxls : les fichiers .xls exigent l'extra urban-json-excel[xls] (xlrd)") from exc
        book = xlrd.open_workbook(path)
        sheets = {}
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append(bool(cell.value))
                    else:
                        row.append(cell.value)
                rows.append(tuple(row))
            sheets[sheet.name] = rows
        return sheets
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in workbook.worksheets}
    finally:
        workbook.close()


def _label(name: str) -> str:
    """Assainit un nom de colonne (``ExcelReader.GetLabel``, comme ``loadcsv``)."""
    decomposed = unicodedata.normalize("NFKD", name)
    out = []
    for char in decomposed:
        if unicodedata.combining(char):
            continue
        if char.isascii() and char.isalnum():
            out.append(char)
        elif char in (" ", "'", '"', "_"):
            out.append("_")
    label = "".join(out)
    while "__" in label:
        label = label.replace("__", "_")
    return label


def _is_integral(value: Any) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, int):
        return True
    return isinstance(value, float) and value.is_integer()


def _auto_columns(header: list[Any], data: list[tuple], header_line: int) -> list[_Column]:
    columns: list[_Column] = []
    width = max([len(header)] + [len(r) for r in data]) if (header or data) else 0
    if header_line > 0:
        names = data[header_line] if header_line < len(data) else ()
        for index in range(width):
            name = names[index] if index < len(names) else None
            if isinstance(name, str):
                text = name.strip(" \t\r\n")
                label = _label(text) if text else f"renamed_{len(columns)}"
            elif name is not None:
                label = f"renamed_{len(columns)}"
            else:
                continue  # colonne sans en-tête : ignorée, comme le .NET
            columns.append(_Column(index, label, "auto"))
        return columns
    seen: set[str] = set()
    for index in range(width):
        raw = header[index] if index < len(header) and header[index] not in (None, "") else f"Column{index}"
        name = _label(str(raw)) or f"Column{index}"
        if name in seen:
            name += f"_renamed_{index}"
        seen.add(name)
        values = [r[index] for r in data if index < len(r) and r[index] is not None]
        kind = "integer" if values and all(_is_integral(v) for v in values) else "auto"
        columns.append(_Column(index, name, kind))
    return columns


def _iso(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _convert(value: Any, column: _Column) -> Any:
    """Valeur typée d'une cellule, ou ``_SKIP`` pour l'omettre."""
    if value is None:
        return _SKIP
    # Colonne automatique contenant du texte : traitée comme une colonne texte (le
    # .NET lui attribue le type string, donc nettoyage et valeurs vides omises).
    if column.type == "string" or (column.type == "auto" and isinstance(value, str)):
        text = value if isinstance(value, str) else str(_iso(value))
        if column.trimmed:
            text = text.strip()
        return _SKIP if column.bypass_empty and text == "" else text
    if isinstance(value, str) and value.strip() == "":
        return _SKIP
    try:
        if column.type == "integer":
            return int(float(value)) if not isinstance(value, bool) else int(value)
        if column.type == "number":
            return float(value)
        if column.type == "boolean":
            return value if isinstance(value, bool) else str(value).strip().lower() in ("true", "1", "vrai", "yes")
    except (TypeError, ValueError):
        return _iso(value)  # conversion impossible : valeur d'origine, comme le .NET
    if column.type == "datetime" and isinstance(value, str):
        return value
    return _iso(value)


_SKIP = object()


def _read(path: str, config: dict) -> list[dict]:
    if not os.path.isfile(path):
        raise JsltError(f"loadxls : classeur introuvable {path!r}")
    sheets = _read_sheets(path)
    table = _get(config, "Tablename") or ""
    if table:
        if table not in sheets:
            raise JsltError(f"loadxls : feuille {table!r} introuvable (feuilles : {', '.join(sheets)})")
    elif len(sheets) == 1:
        table = next(iter(sheets))
    else:
        raise JsltError("loadxls : Tablename non précisé alors que le classeur a plusieurs feuilles")

    rows = sheets[table]
    first_row_names = bool(_get(config, "FirstRowHasColumnName", False))
    header = list(rows[0]) if first_row_names and rows else []
    data = rows[1:] if first_row_names else rows
    header_line = int(_get(config, "HeaderLine", 0) or 0)
    add_row_number = bool(_get(config, "AddRowNumber", True))
    bypass_empty_object = bool(_get(config, "ByPassEmptyObject", True))
    columns = _columns_from_config(config) or _auto_columns(header, data, header_line)

    result = []
    for row_line, row in enumerate(data):
        if header_line and row_line <= header_line:
            continue
        obj: dict[str, Any] = {}
        for column in columns:
            value = row[column.index] if column.index < len(row) else None
            converted = _convert(value, column)
            if converted is not _SKIP:
                obj[column.name] = converted
        if obj or not bypass_empty_object:
            if add_row_number:
                obj["$row"] = row_line
            result.append(obj)
    return result


def execute_loadxlssheet(ctx: Any, source_path: Any = None, config: Any = None) -> list:
    if not source_path:
        raise JsltError("loadxlssheet : le chemin du classeur est requis")
    if config is None:
        config = {}
    if not isinstance(config, dict):
        raise JsltError("loadxlssheet : la configuration doit être un objet JSON")
    return _read(str(source_path), config)


def execute_loadxls(ctx: Any, source_path: Any = None, config_path: Any = None) -> list:
    if not source_path:
        raise JsltError("loadxls : le chemin du classeur est requis")
    if not config_path or not os.path.isfile(str(config_path)):
        raise JsltError(f"loadxls : fichier de configuration introuvable {config_path!r}")
    try:
        with open(str(config_path), encoding="utf-8-sig") as handle:
            config = json.load(handle)
    except json.JSONDecodeError as exc:
        raise JsltError(f"loadxls : fichier de configuration JSON invalide ({exc})") from exc
    if not isinstance(config, dict):
        raise JsltError("loadxls : la configuration doit être un objet JSON")
    return _read(str(source_path), config)


def register(registry: Any) -> None:
    """Enregistre ``loadxls`` et ``loadxlssheet``."""
    registry.register(
        "loadxls", execute_loadxls,
        description="Lit une feuille Excel en tableau d'objets, selon un fichier de configuration JSON",
    )
    registry.register(
        "loadxlssheet", execute_loadxlssheet,
        description="Lit une feuille Excel en tableau d'objets, selon un objet de configuration",
    )
